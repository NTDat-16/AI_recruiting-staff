import io
import os
import sys
sys.stdout.reconfigure(encoding="utf-8")
import asyncio
from PIL import Image
from sqlalchemy import text
from app.core.database import async_engine
from app.modules.candidate.cv_parser import CVParser

def generate_sample_avatar_pdf() -> bytes:
    """Tạo 1 file PDF CV chuẩn có kèm ảnh chân dung avatar ứng viên."""
    # Tạo ảnh chân dung 240x320 (chuẩn tỷ lệ 3:4 chân dung)
    avt = Image.new("RGB", (240, 320), color=(52, 152, 219))
    # Vẽ vài chi tiết đơn giản lên ảnh để giống ảnh thật
    from PIL import ImageDraw
    draw = ImageDraw.Draw(avt)
    # Vẽ đầu và vai
    draw.ellipse((70, 40, 170, 140), fill=(245, 205, 175)) # Đầu
    draw.rectangle((40, 150, 200, 310), fill=(41, 128, 185)) # Áo
    draw.ellipse((85, 75, 105, 95), fill=(40, 40, 40)) # Mắt trái
    draw.ellipse((135, 75, 155, 95), fill=(40, 40, 40)) # Mắt phải
    draw.arc((95, 100, 145, 125), start=0, end=180, fill=(180, 50, 50), width=3) # Miệng cười

    buf = io.BytesIO()
    avt.save(buf, format="PDF", resolution=100.0)
    return buf.getvalue()

def generate_sample_avatar_docx() -> bytes:
    """Tạo 1 file DOCX CV có đính kèm ảnh chân dung."""
    import docx
    doc = docx.Document()
    doc.add_heading("CV ỨNG VIÊN - KỸ SƯ AI", 0)
    doc.add_paragraph("Họ và tên: Lê Văn Bảo\nEmail: levanbao.ai@gmail.com\nSĐT: 0988776655")
    
    # Tạo ảnh chân dung
    avt = Image.new("RGB", (200, 260), color=(46, 204, 113))
    from PIL import ImageDraw
    draw = ImageDraw.Draw(avt)
    draw.ellipse((60, 30, 140, 110), fill=(255, 224, 189))
    draw.rectangle((30, 120, 170, 250), fill=(39, 174, 96))
    
    temp_img = "temp_docx_avt.jpg"
    avt.save(temp_img, "JPEG")
    doc.add_picture(temp_img, width=docx.shared.Inches(1.5))
    
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    if os.path.exists(temp_img):
        try: os.remove(temp_img)
        except: pass
    return doc_buf.getvalue()

def test_unit_extraction():
    print("\n--- 1. KIỂM THỬ TRÍCH XUẤT AVATAR TỪ PDF & DOCX ---")
    pdf_bytes = generate_sample_avatar_pdf()
    pdf_avt = CVParser.extract_avatar_from_bytes(pdf_bytes, "cv_avatar_test.pdf")
    assert pdf_avt is not None, "FAILED: Không trích xuất được avatar từ PDF!"
    print(f"✓ Trích xuất avatar từ PDF thành công! Dung lượng ảnh: {len(pdf_avt)} bytes")

    # Kiểm tra mở ảnh trích xuất bằng PIL
    img_from_pdf = Image.open(io.BytesIO(pdf_avt))
    print(f"✓ Thông số ảnh PDF trích xuất: Định dạng={img_from_pdf.format}, Kích thước={img_from_pdf.size}")

    docx_bytes = generate_sample_avatar_docx()
    docx_avt = CVParser.extract_avatar_from_bytes(docx_bytes, "cv_avatar_test.docx")
    assert docx_avt is not None, "FAILED: Không trích xuất được avatar từ DOCX!"
    print(f"✓ Trích xuất avatar từ DOCX thành công! Dung lượng ảnh: {len(docx_avt)} bytes")
    img_from_docx = Image.open(io.BytesIO(docx_avt))
    print(f"✓ Thông số ảnh DOCX trích xuất: Định dạng={img_from_docx.format}, Kích thước={img_from_docx.size}")

    # Test file không có ảnh (Plain text)
    txt_bytes = b"Ho va ten: Tran Van A\nEmail: a@test.com"
    txt_avt = CVParser.extract_avatar_from_bytes(txt_bytes, "cv_plain.txt")
    assert txt_avt is None, "FAILED: File plain text không được có avatar!"
    print("✓ File không có ảnh (plain text) trả về None chuẩn xác, không ném exception.")

if __name__ == "__main__":
    test_unit_extraction()
