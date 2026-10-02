import io
import logging
from typing import Tuple, Optional
from app.ai.llm_client import get_llm_client
from app.ai.schemas import ParsedCVSchema

logger = logging.getLogger(__name__)


class CVParser:
    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        """Trích xuất văn bản thô từ file PDF hoặc DOCX."""
        text = ""
        filename_lower = filename.lower()
        try:
            if filename_lower.endswith(".pdf"):
                try:
                    import pdfplumber
                    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                        pages = [page.extract_text() or "" for page in pdf.pages]
                        text = "\n".join(pages)
                except Exception as e:
                    logger.warning(f"pdfplumber extraction failed: {e}")
                    text = file_bytes.decode("utf-8", errors="ignore")
            elif filename_lower.endswith(".docx"):
                try:
                    import docx
                    doc = docx.Document(io.BytesIO(file_bytes))
                    text = "\n".join([p.text for p in doc.paragraphs if p.text])
                except Exception as e:
                    logger.warning(f"docx extraction failed: {e}")
                    text = file_bytes.decode("utf-8", errors="ignore")
            else:
                # Text/markdown/etc
                text = file_bytes.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"Error parsing file {filename}: {e}")
            text = file_bytes.decode("utf-8", errors="ignore")

        return text.strip()

    @staticmethod
    def extract_avatar_from_bytes(file_bytes: bytes, filename: str) -> Optional[bytes]:
        """Trích xuất ảnh đại diện (avatar) của ứng viên từ file PDF hoặc DOCX nếu có."""
        if not file_bytes or not filename:
            return None

        filename_lower = filename.lower()
        try:
            from PIL import Image

            if filename_lower.endswith(".pdf"):
                try:
                    import pypdfium2
                    doc = pypdfium2.PdfDocument(file_bytes)
                    candidate_images = []

                    # Ứng viên thường đặt ảnh chân dung ở Trang 1 (hoặc Trang 2)
                    pages_to_check = [0] if len(doc) > 0 else []
                    if len(doc) > 1:
                        pages_to_check.append(1)

                    for page_idx in pages_to_check:
                        page = doc[page_idx]
                        page_w = page.get_width()
                        page_h = page.get_height()

                        for obj in page.get_objects():
                            if obj.type == pypdfium2.raw.FPDF_PAGEOBJ_IMAGE:
                                try:
                                    bm = obj.get_bitmap()
                                    img = bm.to_pil()
                                    w, h = img.size

                                    # 1. Bỏ qua icon nhỏ / bullet / logo li ti
                                    if w < 50 or h < 50 or (w * h) < 3600:
                                        continue

                                    # 2. Bỏ qua trang scan toàn bộ (full-page scan có kích thước lớn của cả trang A4)
                                    bounds = obj.get_bounds()
                                    box_h = bounds[3] - bounds[1]
                                    box_w = bounds[2] - bounds[0]
                                    if page_h > 0 and (box_h / page_h) > 0.85 and (box_w / page_w) > 0.85 and (w >= 1000 and h >= 1200):
                                        continue

                                    # 3. Tỷ lệ ảnh: Ảnh đại diện/chân dung thường là hình chữ nhật đứng (3x4, 4x6) hoặc vuông (1:1)
                                    aspect = w / h
                                    if aspect < 0.5 or aspect > 1.8:
                                        continue

                                    # 4. Chấm điểm vị trí: Ảnh đại diện thường nằm ở nửa trên của trang
                                    pos_y = bounds[3] / page_h if page_h > 0 else 0.5
                                    page_weight = 1.0 if page_idx == 0 else 0.6

                                    aspect_score = 1.0 - min(abs(aspect - 0.85), 1.0)
                                    position_score = pos_y
                                    size_score = min(w, h) / 400.0 if min(w, h) <= 400 else 400.0 / min(w, h)

                                    total_score = ((aspect_score * 0.4) + (position_score * 0.4) + (size_score * 0.2)) * page_weight
                                    candidate_images.append((total_score, img))
                                except Exception as err:
                                    logger.debug(f"Bỏ qua đối tượng ảnh PDF không hợp lệ: {err}")

                    doc.close()

                    if candidate_images:
                        candidate_images.sort(key=lambda x: x[0], reverse=True)
                        best_img = candidate_images[0][1]

                        if best_img.mode in ("RGBA", "P"):
                            best_img = best_img.convert("RGB")
                        elif best_img.mode != "RGB":
                            best_img = best_img.convert("RGB")

                        best_img.thumbnail((400, 400), Image.Resampling.LANCZOS)
                        out_buf = io.BytesIO()
                        best_img.save(out_buf, format="JPEG", quality=85)
                        return out_buf.getvalue()

                except Exception as e:
                    logger.warning(f"Lỗi khi trích xuất avatar từ PDF: {e}")

            elif filename_lower.endswith(".docx"):
                try:
                    import docx
                    doc = docx.Document(io.BytesIO(file_bytes))
                    candidate_images = []

                    for rel in doc.part.related_parts.values():
                        if "image" in getattr(rel, "content_type", ""):
                            try:
                                img = Image.open(io.BytesIO(rel.blob))
                                w, h = img.size

                                if w < 50 or h < 50 or (w * h) < 3600:
                                    continue

                                aspect = w / h
                                if aspect < 0.5 or aspect > 1.8:
                                    continue

                                aspect_score = 1.0 - min(abs(aspect - 0.85), 1.0)
                                size_score = min(w, h) / 400.0 if min(w, h) <= 400 else 400.0 / min(w, h)
                                score = aspect_score * 0.6 + size_score * 0.4
                                candidate_images.append((score, img))
                            except Exception as err:
                                logger.debug(f"Bỏ qua ảnh docx: {err}")

                    if candidate_images:
                        candidate_images.sort(key=lambda x: x[0], reverse=True)
                        best_img = candidate_images[0][1]

                        if best_img.mode in ("RGBA", "P"):
                            best_img = best_img.convert("RGB")
                        elif best_img.mode != "RGB":
                            best_img = best_img.convert("RGB")

                        best_img.thumbnail((400, 400), Image.Resampling.LANCZOS)
                        out_buf = io.BytesIO()
                        best_img.save(out_buf, format="JPEG", quality=85)
                        return out_buf.getvalue()

                except Exception as e:
                    logger.warning(f"Lỗi khi trích xuất avatar từ DOCX: {e}")

        except Exception as e:
            logger.warning(f"Lỗi tổng quát khi trích xuất avatar: {e}")

        return None

    @staticmethod
    async def parse_and_structure_cv(file_bytes: bytes, filename: str) -> Tuple[str, ParsedCVSchema]:
        """Trích xuất văn bản và gọi AI layer để chuẩn hóa thông tin ứng viên."""
        raw_text = CVParser.extract_text_from_bytes(file_bytes, filename)
        llm = get_llm_client()
        parsed_schema = await llm.parse_cv_text(raw_text)
        return raw_text, parsed_schema
